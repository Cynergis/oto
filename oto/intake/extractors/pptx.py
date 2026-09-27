# -*- coding: utf-8 -*-
"""PowerPoint to a Document: title, body text, tables, speaker notes, and figure detection.

A deck carries much of its meaning in pictures, so each slide is classified. A slide counts as a
figure when it holds a chart, a grouped vector graphic, several connectors, a dominant picture, or
three or more pictures. Those signals mean "the words on this slide do not carry its content", which
is exactly when a description is needed for the slide to be searchable.

`level` on a bullet block is the outline indent, not a heading depth.
"""
import os

from ..document import BULLET, DIAGRAM, NOTE, PAGE, TABLE, Block, Document, clean_cell, slugify, \
    table_to_markdown
from ..registry import Extractor, register

DOMINANT_PICTURE_AREA = 0.20     # share of the slide a single picture must cover to dominate
MANY_PICTURES = 3
MANY_CONNECTORS = 2


def _shape_area(shape):
    try:
        return (shape.width or 0) * (shape.height or 0)
    except Exception:
        return 0


def _paragraphs(shape):
    """Outline level and text for each non-empty paragraph in a shape."""
    out = []
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            text = "".join(r.text for r in p.runs) or p.text
            if text.strip():
                out.append((p.level or 0, text.strip()))
    return out


def _classify(slide, slide_area):
    """Return (is_figure, reasons, counts). Reason order is part of the rendered output."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    counts = {"pictures": 0, "charts": 0, "groups": 0, "connectors": 0}
    dominant_picture = [False]

    def walk(shapes):
        for shape in shapes:
            kind = shape.shape_type
            if kind == MSO_SHAPE_TYPE.PICTURE:
                counts["pictures"] += 1
                if slide_area and _shape_area(shape) > DOMINANT_PICTURE_AREA * slide_area:
                    dominant_picture[0] = True
            elif kind == MSO_SHAPE_TYPE.CHART:
                counts["charts"] += 1
            elif kind == MSO_SHAPE_TYPE.GROUP:
                counts["groups"] += 1
                try:
                    walk(shape.shapes)
                except Exception:
                    pass
            elif kind == MSO_SHAPE_TYPE.LINE or (kind is not None and "CONNECTOR" in str(kind)):
                counts["connectors"] += 1

    walk(slide.shapes)

    reasons = []
    if counts["charts"]:
        reasons.append("%d chart(s)" % counts["charts"])
    if counts["groups"]:
        reasons.append("%d grouped graphic(s)" % counts["groups"])
    if counts["connectors"] >= MANY_CONNECTORS:
        reasons.append("%d connector(s)/flow lines" % counts["connectors"])
    if dominant_picture[0]:
        reasons.append("dominant picture")
    if counts["pictures"] >= MANY_PICTURES:
        reasons.append("%d pictures" % counts["pictures"])
    return bool(reasons), reasons, counts


def _title(slide):
    """The title placeholder, or the highest non-empty text box as a fallback."""
    try:
        if slide.shapes.title and slide.shapes.title.text.strip():
            return slide.shapes.title.text.strip()
    except Exception:
        pass
    best, best_top = None, None
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            top = shape.top or 0
            if best is None or top < best_top:
                best = shape.text_frame.text.strip().split("\n")[0]
                best_top = top
    return best or ""


@register
class Pptx(Extractor):
    name = "pptx"
    extensions = (".pptx",)
    media_type = "PowerPoint deck"
    requires = ("pptx",)

    def extract(self, path, style, assets_dir=None):
        from pptx import Presentation

        filename = os.path.basename(path)
        slug = slugify(filename)
        doc = Document(slug=slug, title=os.path.splitext(filename)[0],
                       source_name=filename, media_type=self.media_type)
        if assets_dir:
            os.makedirs(assets_dir, exist_ok=True)

        deck = Presentation(path)
        slide_area = (deck.slide_width or 0) * (deck.slide_height or 0)
        n_slides = 0

        for number, slide in enumerate(deck.slides, 1):
            n_slides = number
            title = _title(slide)
            is_figure, reasons, counts = _classify(slide, slide_area)

            title_shape = None
            try:
                title_shape = slide.shapes.title
            except Exception:
                pass

            body, tables = [], []
            for shape in slide.shapes:
                if shape is title_shape:
                    continue
                if shape.has_table:
                    tables.append(shape.table)
                    continue
                body.extend(_paragraphs(shape))

            notes = ""
            try:
                if slide.has_notes_slide:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
            except Exception:
                pass

            # `had_body` records that the slide had text before de-duplication. The renderer emits a
            # blank line based on that, so it must survive as a fact about the slide.
            doc.add(Block(PAGE, text=title, level=number,
                          meta={"had_body": bool(body), "counts": counts}))
            if is_figure:
                asset = "assets/%s/slide-%02d.png" % (slug, number)
                doc.assets.append(asset)
                doc.add(Block(DIAGRAM, path=asset, meta={"slide": number, "reasons": reasons}))

            seen = set()
            for level, text in body:
                if (level, text) in seen or text == title:
                    seen.add((level, text))
                    continue
                seen.add((level, text))
                doc.add(Block(BULLET, text=text, level=level))

            for table in tables:
                rows = [[clean_cell(c.text) for c in r.cells] for r in table.rows]
                if rows:
                    doc.add(Block(TABLE, rows=rows))

            if notes:
                doc.add(Block(NOTE, text=notes))

        doc.extra = {"n_slides": n_slides}
        return doc


def render(doc, style):
    """Render to markdown. Byte-compatible with the legacy pptx extractor."""
    header = ("# %s\n\n> **Source:** `%s` · **Type:** %s · **Slides:** %d · Converted to Markdown "
              "by the %s extraction pipeline.\n"
              % (doc.title, style.ref(doc.source_name), doc.media_type,
                 doc.extra.get("n_slides", 0), style.pipeline_name))
    lines = []
    pending_body_blank = False

    def flush_body_blank():
        if pending_body_blank:
            lines.append("")

    for b in doc.blocks:
        if b.kind == PAGE:
            flush_body_blank()
            pending_body_blank = False
            suffix = (" — " + b.text) if b.text else ""
            lines.append("\n---\n\n## Slide %d%s\n" % (b.level, suffix))
            pending_body_blank = b.meta.get("had_body", False)
        elif b.kind == DIAGRAM:
            lines.append("> **[DIAGRAM]** (%s) — see `%s`. _Description added in review pass._\n"
                         % (", ".join(b.meta["reasons"]), b.path))
            lines.append("![Slide %d diagram](%s)\n" % (b.meta["slide"], b.path))
        elif b.kind == BULLET:
            lines.append(("  " * b.level) + "- %s" % b.text)
        elif b.kind == TABLE:
            flush_body_blank()
            pending_body_blank = False
            md = table_to_markdown(b.rows)
            if md:
                lines.append(md + "\n")
        elif b.kind == NOTE:
            flush_body_blank()
            pending_body_blank = False
            lines.append("\n**Speaker notes:** %s\n" % b.text)

    flush_body_blank()
    return header + "\n".join(lines) + "\n"

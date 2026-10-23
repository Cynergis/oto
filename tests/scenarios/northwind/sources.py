"""The documents Northwind Transit's analyst starts with, written at test time: a policy (.docx), a stock
sheet (.xlsx), an incident review and the telematics notes. The notes carry a credential, so the ingest
must block them; the key is assembled here rather than written whole, so no scanner reads the
repository as leaking one."""
import os

KEY = "AKIA" + "IOSFODNN7EXAMPLE"   # AWS's documented example key: inert, and shaped like a real one


def make(out):
    from docx import Document
    from openpyxl import Workbook

    os.makedirs(out, exist_ok=True)
    d = Document()
    d.add_heading("Fleet Maintenance Policy 2025", 0)
    d.add_heading("Purpose", 1)
    d.add_paragraph("This policy sets how Northwind Transit keeps its bus fleet safe and in service. "
                    "It applies to every Bus in the fleet and to every Garage that maintains one.")
    d.add_heading("Roles", 1)
    d.add_paragraph("The Fleet Manager owns this policy and approves every change to the PM Schedule.")
    d.add_paragraph("The Maintenance Supervisor at each Garage assigns Work Orders to technicians "
                    "and closes them once the Inspection passes.")
    d.add_paragraph("The Safety Officer reviews every Incident Report within five working days.")
    d.add_heading("Preventive Maintenance", 1)
    d.add_paragraph("Every Bus follows the PM Schedule. A missed PM interval opens a Work Order automatically.")
    t = d.add_table(rows=1, cols=3)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text, t.rows[0].cells[2].text = "PM level", "Interval", "Performed at"
    for row in (("PM-A", "every 6,000 km", "any Garage"), ("PM-B", "every 24,000 km", "Central Garage"),
                ("PM-C", "every 96,000 km", "Central Garage")):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    d.add_heading("Work Orders", 1)
    d.add_paragraph("A Work Order names the Bus, the Garage, the fault and the parts used. "
                    "A Work Order raised after an Incident Report cites it.")
    d.add_paragraph("Effective 2025-01-15. Supersedes the 2022 maintenance procedure.")
    d.save(os.path.join(out, "Fleet Maintenance Policy 2025.docx"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Q3 stock"
    ws.append(["Part number", "Description", "Garage", "On hand", "Reorder at"])
    for row in (("BRK-220", "Brake pad set, front", "Central Garage", 40, 25),
                ("BRK-221", "Brake pad set, rear", "Central Garage", 18, 25),
                ("FLT-010", "Oil filter", "East Garage", 120, 60),
                ("TYR-305", "Tyre 305/70R22.5", "East Garage", 12, 16)):
        ws.append(list(row))
    wb.save(os.path.join(out, "Parts Inventory Q3.xlsx"))

    with open(os.path.join(out, "Incident Review - Bus 4127.md"), "w", encoding="utf-8") as f:
        f.write("""# Incident Review: Bus 4127 brake fault

Date of incident: 2025-08-02. Reviewed by the Safety Officer on 2025-08-05.

## What happened

Bus 4127 (VIN 1HGBH41JXMN109186) reported soft brakes on Route 12. The driver returned it to
East Garage. The Maintenance Supervisor opened Work Order WO-2291 the same day.

## Findings

- Rear brake pads (BRK-221) were below the wear limit.
- The bus was 1,800 km past its PM-B interval: the PM Schedule alert did not fire.
- East Garage had no BRK-221 in stock; parts came from Central Garage.

## Actions

- Replace rear pads on Bus 4127 under WO-2291.
- Fleet Manager to review why the PM-B alert failed.
- Questions to the review lead: s.okafor@northwind-transit.example
""")

    with open(os.path.join(out, "Vendor Integration Notes.txt"), "w", encoding="utf-8") as f:
        f.write("Telematics vendor integration notes\n\n"
                "The telematics feed posts odometer readings for every Bus nightly.\n"
                "Test credentials used during setup: %s\n"
                "The PM Schedule alert reads the odometer from this feed.\n" % KEY)
